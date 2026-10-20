import io
import os
import openpyxl
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="設備投資計画フォーム 自動集約ツール",
    page_icon="📊",
    layout="wide"
)

st.title("📊 設備投資計画フォーム 自動集約ツール")
st.markdown("""
各部署から回収した「設備投資計画フォーム（Excel）」をまとめてドラッグ＆ドロップしてください。
元の「全国まとめフォーマット（計算式やデザイン込み）」に各部署のデータを自動で集約・転記します。
""")

# シートを特定するためのキーワードマッピング
TARGET_SHEET_KEYWORDS = [
    ("新規取得_年初", ["①", "新規取得", "年初"]),
    ("新規取得_追加", ["①", "新規取得", "追加用"]),
    ("補修リース_年初", ["②", "補修", "年初"]),
    ("補修リース_追加", ["②", "補修", "追加用"]),
    ("前年度追加", ["⓷", "新規取得", "追加"]),
]


def find_matching_sheet(sheet_list, keywords):
    """シート名一覧からキーワードに合致するシート名を検索"""
    for s in sheet_list:
        if all(kw in s for kw in keywords):
            return s
    return None


st.sidebar.header("📂 テンプレート設定")
template_file = st.sidebar.file_uploader(
    "【任意】まとめ用テンプレート（0000年度/全国まとめフォーム）",
    type=["xlsx"]
)

uploaded_files = st.file_uploader(
    "各部署のExcelファイルをここにドラッグ＆ドロップ（複数可）",
    type=["xlsx"],
    accept_multiple_files=True
)

if uploaded_files:
    st.info(f"📁 現在 {len(uploaded_files)} 個のファイルが選択されています。")

    if st.button("⚡ 集計処理を実行する", type="primary"):
        combined_rows = {key: [] for key, _ in TARGET_SHEET_KEYWORDS}
        processed_count = 0

        # 各部署ファイルからデータ抽出
        for uploaded_file in uploaded_files:
            filename = uploaded_file.name
            if "★" in filename or "全国" in filename or "集約" in filename or "0000年度" in filename:
                continue

            try:
                xl = pd.ExcelFile(uploaded_file)
                for key, keywords in TARGET_SHEET_KEYWORDS:
                    matched_sheet = find_matching_sheet(xl.sheet_names, keywords)
                    if matched_sheet:
                        df = pd.read_excel(uploaded_file, sheet_name=matched_sheet, header=7)
                        if "資産名称" in df.columns:
                            df_valid = df.dropna(subset=["資産名称"]).copy()
                            if "No." in df_valid.columns:
                                df_valid = df_valid[df_valid["No."].astype(str) != "合計"]

                            if not df_valid.empty:
                                combined_rows[key].append(df_valid)

                processed_count += 1
            except Exception as e:
                st.error(f"❌ {filename} の読み込みエラー: {e}")

        # テンプレートファイルの選択
        base_file = template_file if template_file else uploaded_files[0]

        try:
            wb = openpyxl.load_workbook(base_file)

            for key, keywords in TARGET_SHEET_KEYWORDS:
                matched_sheet = find_matching_sheet(wb.sheetnames, keywords)
                if matched_sheet and combined_rows[key]:
                    ws = wb[matched_sheet]

                    all_data_df = pd.concat(combined_rows[key], ignore_index=True)

                    # No.を1からの連番にする
                    if "No." in all_data_df.columns:
                        all_data_df["No."] = range(1, len(all_data_df) + 1)

                    # 9行目から順にデータを転記
                    start_row = 9
                    for r_idx, row in all_data_df.iterrows():
                        current_row = start_row + r_idx
                        for c_idx, val in enumerate(row):
                            if pd.notna(val):
                                ws.cell(row=current_row, column=c_idx + 1, value=val)

            output = io.BytesIO()
            wb.save(output)

            st.success(f"🎉 {processed_count} 個のファイルから集約・転記が完了しました！")

            st.download_button(
                label="📥 集約版Excelファイルをダウンロード",
                data=output.getvalue(),
                file_name="★【全国】設備投資計画集約版.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

        except Exception as e:
            st.error(f"集約ファイルの生成中にエラーが発生しました: {e}")
